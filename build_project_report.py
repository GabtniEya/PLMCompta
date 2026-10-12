from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path
from datetime import date

OUT = Path('output') / 'Rapport_Projet_PLM_Ressource.docx'
NAVY = '0B2545'; TEAL = '0F8B82'; PALE = 'E8F4F2'; LIGHT = 'F2F4F7'; MUTED = '64748B'

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def margins(cell):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); m = tcPr.first_child_found_in('w:tcMar')
    if m is None: m = OxmlElement('w:tcMar'); tcPr.append(m)
    for side in ('top','start','bottom','end'):
        e = OxmlElement(f'w:{side}'); e.set(qn('w:w'), '120'); e.set(qn('w:type'), 'dxa'); m.append(e)

def set_cell_text(cell, text, bold=False, color=NAVY, size=10.5):
    cell.text = ''
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(text)); r.bold = bold; r.font.name = 'Calibri'; r._element.rPr.rFonts.set(qn('w:ascii'),'Calibri'); r.font.size = Pt(size); r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; margins(cell)

def set_table_width(table, widths):
    table.autofit = False
    for row in table.rows:
        for i, width in enumerate(widths): row.cells[i].width = Inches(width)

def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}'); p.add_run(text); return p

def add_bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet'); p.add_run(text); return p

def add_number(doc, text):
    p = doc.add_paragraph(style='List Number'); p.add_run(text); return p

doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = sec.footer_distance = Inches(.492)

normal = doc.styles['Normal']; normal.font.name = 'Calibri'; normal._element.rPr.rFonts.set(qn('w:ascii'),'Calibri'); normal.font.size = Pt(11); normal.font.color.rgb = RGBColor.from_string('1F2937'); normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.1
for name, size, color, before, after in [('Heading 1',16,'2E74B5',16,8),('Heading 2',13,'2E74B5',12,6),('Heading 3',12,'1F4D78',8,4)]:
    st=doc.styles[name]; st.font.name='Calibri'; st._element.rPr.rFonts.set(qn('w:ascii'),'Calibri'); st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(color); st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after)

header = sec.header.paragraphs[0]; header.alignment = WD_ALIGN_PARAGRAPH.RIGHT; r = header.add_run('PLM RESSOURCE  |  RAPPORT PROJET'); r.font.size=Pt(8.5); r.font.color.rgb=RGBColor.from_string(MUTED)
footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = footer.add_run('Document de synthèse - version de démonstration'); r.font.size=Pt(8.5); r.font.color.rgb=RGBColor.from_string(MUTED)

# Cover
doc.add_paragraph().paragraph_format.space_after = Pt(72)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('PLM RESSOURCE'); r.bold=True; r.font.name='Calibri'; r.font.size=Pt(14); r.font.color.rgb=RGBColor.from_string(TEAL)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('Rapport de projet'); r.bold=True; r.font.name='Calibri'; r.font.size=Pt(30); r.font.color.rgb=RGBColor.from_string(NAVY)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('Plateforme de services professionnels, commandes et administration'); r.font.size=Pt(14); r.font.color.rgb=RGBColor.from_string(MUTED)
doc.add_paragraph().paragraph_format.space_after = Pt(58)
t=doc.add_table(rows=4, cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'; set_table_width(t,[1.8,4.7])
for row,(a,b) in zip(t.rows,[('Projet','PLM Ressource'),('Périmètre','Vitrine, commande, facturation et administration'),('Environnement','Application web React + Supabase'),('Date',date.today().strftime('%d/%m/%Y'))]):
    shade(row.cells[0],PALE); set_cell_text(row.cells[0],a,True,TEAL); set_cell_text(row.cells[1],b)
doc.add_paragraph().paragraph_format.space_after=Pt(40)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('Statut : prêt pour démonstration et tests fonctionnels'); r.bold=True; r.font.size=Pt(11); r.font.color.rgb=RGBColor.from_string(TEAL)
doc.add_page_break()

add_heading(doc,'1. Résumé exécutif')
doc.add_paragraph("PLM Ressource est une application web destinée à présenter des services professionnels en Tunisie, permettre à un client de préparer une commande et offrir un espace d'administration pour le suivi commercial. Le projet couvre le parcours public de la vitrine ainsi qu'un espace de gestion des clients, factures et paiements.")
callout=doc.add_table(rows=1,cols=1); shade(callout.cell(0,0),PALE); set_cell_text(callout.cell(0,0),'Point clé : le paiement est volontairement en mode test. Aucun débit de carte bancaire ni passerelle réelle n’est activé dans cette version.',True,TEAL,10.5)
add_heading(doc,'Objectifs atteints',2)
for x in ['Présenter un catalogue de prestations avec prix en dinars tunisiens.','Permettre l’ajout au panier et la création d’une commande.','Générer un numéro de commande et une facture PDF téléchargeable.','Centraliser le suivi commercial dans une interface d’administration.','Prévoir une architecture Supabase sécurisée pour le passage en production.'] : add_bullet(doc,x)

add_heading(doc,'2. Fonctionnalités livrées')
t=doc.add_table(rows=1, cols=3); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER; set_table_width(t,[2.0,3.35,1.15])
for c,tx in zip(t.rows[0].cells,['Module','Fonctionnalités','Statut']): shade(c,LIGHT); set_cell_text(c,tx,True,NAVY)
rows=[('Vitrine publique','Accueil, présentation, services, prix et coordonnées.','Actif'),('Panier et commande','Quantités, total, coordonnées client et numéro de commande.','Actif'),('Facture PDF','Récapitulatif client, lignes de prestation, total et statut.','Actif'),('Administration','Connexion, tableau de bord, clients, factures, paiements, services.','Actif'),('Paiement','Formulaire de carte simulé pour les tests uniquement.','Mode test'),('Paiement réel','Passerelle de paiement en ligne.','Non activé')]
for a,b,c in rows:
    cells=t.add_row().cells; set_cell_text(cells[0],a,True); set_cell_text(cells[1],b); set_cell_text(cells[2],c,True,TEAL if c=='Actif' else '7A5A00')

add_heading(doc,'3. Parcours utilisateur et tests',1)
add_heading(doc,'Test du client',2)
for x in ['Ouvrir la page d’accueil et parcourir le catalogue.','Ajouter une ou plusieurs prestations au panier.','Ouvrir le panier, puis choisir de passer commande.','Renseigner les coordonnées demandées.','Utiliser la carte de test proposée : 4242 4242 4242 4242, expiration 12/30, CVC 123.','Valider la commande de test, vérifier le numéro de commande et télécharger la facture PDF.'] : add_number(doc,x)
add_heading(doc,'Test de l’administrateur',2)
for x in ['Ouvrir /admin et créer ou utiliser un compte de connexion.','Vérifier le tableau de bord et la liste des services.','Créer un client, puis une facture associée.','Enregistrer un paiement afin de faire passer la facture au statut Payée.','Actualiser la page pour vérifier la persistance après configuration Supabase.'] : add_number(doc,x)

add_heading(doc,'4. Architecture technique',1)
doc.add_paragraph('Le front-end repose sur React, TypeScript et Vite. L’interface utilise Tailwind CSS et Lucide pour les icônes. Supabase fournit l’authentification, la base PostgreSQL et les règles de sécurité. Les factures de commande sont générées dans le navigateur au format PDF.')
t=doc.add_table(rows=1,cols=2); t.style='Table Grid'; set_table_width(t,[1.9,4.6])
for c,tx in zip(t.rows[0].cells,['Composant','Rôle']): shade(c,LIGHT); set_cell_text(c,tx,True)
for a,b in [('React / Vite','Application web, pages publiques et interface d’administration.'),('Supabase Auth','Connexion et gestion de session des utilisateurs.'),('Supabase Database','Services, commandes, clients, factures, paiements et paramètres.'),('RLS','Limitation des accès aux données d’administration.'),('jsPDF','Production de la facture téléchargeable après commande.')]:
    cells=t.add_row().cells; set_cell_text(cells[0],a,True); set_cell_text(cells[1],b)

add_heading(doc,'5. Sécurité et données',1)
for x in ['Aucune vraie donnée de carte ne doit être saisie dans le mode test.','Les clés privées ou API ne doivent jamais être placées dans les variables VITE_ du front-end.','Les politiques RLS Supabase doivent être appliquées avant la publication.','Les comptes administrateurs doivent être explicitement autorisés une fois Supabase configuré.','Le site doit être déployé uniquement derrière HTTPS avec un domaine de production.'] : add_bullet(doc,x)
add_heading(doc,'Limites actuelles',2)
doc.add_paragraph('La version actuelle est une démonstration. Le paiement par carte réel est désactivé à la demande. Les migrations Supabase doivent être exécutées afin d’activer la persistance sécurisée des clients, factures et paiements dans la base distante.')

add_heading(doc,'6. Mise en production : checklist',1)
for x in ['Déployer le site sur un hébergeur compatible HTTPS (Vercel, Netlify ou hébergeur interne).','Connecter le nom de domaine et vérifier le certificat HTTPS.','Exécuter les migrations du dossier supabase/migrations dans Supabase SQL Editor.','Créer le premier compte et lui attribuer le rôle administrateur.','Définir VITE_REQUIRE_ADMIN_ROLE=true sur l’hébergement après le test des rôles.','Renseigner les informations légales de l’entreprise et faire valider le format des factures par le comptable.','Effectuer un test de commande complet et vérifier le PDF téléchargé.'] : add_number(doc,x)

add_heading(doc,'7. Livrables et fichiers utiles',1)
t=doc.add_table(rows=1,cols=2); t.style='Table Grid'; set_table_width(t,[2.3,4.2])
for c,tx in zip(t.rows[0].cells,['Fichier / dossier','Utilité']): shade(c,LIGHT); set_cell_text(c,tx,True)
for a,b in [('src/','Code de l’application React.'),('supabase/migrations/','Schéma et règles de sécurité Supabase.'),('PRODUCTION_SETUP.md','Procédure d’activation de la base, des rôles et du HTTPS.'),('PAYMEE_SETUP.md','Préparation technique existante, non activée dans cette version.'),('package.json','Scripts de lancement, contrôle et compilation.')]:
    cells=t.add_row().cells; set_cell_text(cells[0],a,True); set_cell_text(cells[1],b)

add_heading(doc,'Conclusion',1)
doc.add_paragraph('PLM Ressource dispose d’un parcours cohérent de présentation, commande, facturation et gestion interne. Le projet est prêt pour démonstration. La mise en ligne publique exige l’application des migrations Supabase, la configuration des accès administrateurs et un déploiement HTTPS. Le paiement en ligne réel demeure hors périmètre de cette version.')

OUT.parent.mkdir(parents=True,exist_ok=True)
doc.core_properties.author='PLM Ressource'
doc.core_properties.title='Rapport de projet PLM Ressource'
doc.save(OUT)
print(OUT)
